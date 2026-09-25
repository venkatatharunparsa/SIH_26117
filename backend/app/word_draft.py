"""Word DRAFT generator — proper inspection note with DRAFT badge (G2 / Expected Solution)."""

from __future__ import annotations

import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

from .audit import log_event
from .config import ARTIFACTS_DIR, TEMPLATES_DIR, ensure_dirs

WORD_TEMPLATE = TEMPLATES_DIR / "kwb_inspection_note.docx"


def ensure_base_word_template() -> Path:
    """Seed a plain base .docx template (no logo) if missing."""
    ensure_dirs()
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    if WORD_TEMPLATE.is_file():
        return WORD_TEMPLATE
    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("KWB inspection note base template — sections filled by Workbench")
    r.italic = True
    doc.add_paragraph("{{DRAFT_BANNER}}")
    doc.save(str(WORD_TEMPLATE))
    log_event("word_template_seeded", path=str(WORD_TEMPLATE))
    return WORD_TEMPLATE


def write_inspection_draft(
    *,
    title: str,
    findings: str,
    cites: list[str] | None = None,
    model_id: str | None = None,
    task_id: str | None = None,
    artefact_version: str | None = None,
    grant_id: str | None = None,
    card_id: str | None = None,
    decision: str | None = None,
    asset_id: str | None = None,
    location: str | None = None,
    recommendation: str | None = None,
    actions: str | None = None,
    restrictions: str | None = None,
    monitoring: str | None = None,
    polished: bool = False,
) -> dict:
    """Write a proper DRAFT .docx. Gateway never calls this — Workbench/Orch only."""
    ensure_dirs()
    ensure_base_word_template()
    # Start from base template shell, then rewrite body (keeps file lineage to template)
    try:
        doc = Document(str(WORD_TEMPLATE))
        # Clear seeded paragraphs
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)
    except Exception:
        doc = Document()

    banner = doc.add_paragraph()
    run = banner.add_run("DRAFT — NOT A PLANT RECORD — FOR EXPORT / REVIEW ONLY")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)

    doc.add_heading(title or "Inspection recommendation note", level=1)
    meta = doc.add_paragraph()
    meta.add_run(
        f"Generated: {datetime.now(timezone.utc).isoformat()}  |  "
        f"model_id: {model_id or 'n/a'}  |  task: {task_id or 'n/a'}"
    )
    if card_id:
        meta.add_run(f"  |  card: {card_id}")
    if grant_id:
        meta.add_run(f"  |  grant: {grant_id[:12]}")
    if artefact_version:
        meta.add_run(f"  |  artefact_version: {artefact_version}")
    if polished:
        meta.add_run("  |  body: LLM-polished via Gateway")

    doc.add_heading("Decision / recommendation", level=2)
    doc.add_paragraph((decision or recommendation or "NOT STATED").strip())

    doc.add_heading("Equipment / asset", level=2)
    asset_bits = []
    if asset_id and asset_id.strip():
        asset_bits.append(f"Asset ID: {asset_id.strip()}")
    if location and location.strip():
        asset_bits.append(f"Location: {location.strip()}")
    doc.add_paragraph("\n".join(asset_bits) if asset_bits else "NOT STATED")

    doc.add_heading("Findings / evidence", level=2)
    doc.add_paragraph((findings or "(empty)").strip())

    if recommendation and recommendation.strip() and recommendation.strip() != (decision or "").strip():
        doc.add_heading("Recommendation (detail)", level=2)
        doc.add_paragraph(recommendation.strip())

    doc.add_heading("Actions", level=2)
    doc.add_paragraph((actions or "NOT STATED").strip())

    doc.add_heading("Restrictions", level=2)
    doc.add_paragraph((restrictions or "NOT STATED").strip())

    doc.add_heading("Monitoring", level=2)
    doc.add_paragraph((monitoring or "NOT STATED").strip())

    doc.add_heading("Citations", level=2)
    if cites:
        for c in cites:
            doc.add_paragraph(c, style="List Bullet")
    else:
        doc.add_paragraph("NOT FOUND — no grounded citations")

    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.add_run(
        "Status: DRAFT. KWB assists in the workbench; export leave ends the solution path. "
        "No in-app accept-for-forward. Not a plant record. Not CERT."
    ).italic = True

    name = f"draft-{uuid.uuid4().hex[:10]}.docx"
    path = ARTIFACTS_DIR / name
    doc.save(str(path))
    log_event(
        "word_draft",
        path=str(path),
        model_id=model_id,
        task_id=task_id,
        artefact_version=artefact_version,
        polished=polished,
    )
    return {
        "ok": True,
        "path": str(path),
        "filename": name,
        "status": "DRAFT",
        "format": "docx",
        "artefact_version": artefact_version,
        "polished": polished,
    }


def parse_polish_fields(text: str) -> dict[str, str]:
    """Parse LLM polish output — KEY: value lines (tolerant)."""
    keys = (
        "title",
        "decision",
        "asset_id",
        "location",
        "findings",
        "recommendation",
        "actions",
        "restrictions",
        "monitoring",
    )
    aliases = {
        "title": ("title",),
        "decision": ("decision",),
        "asset_id": ("asset_id", "asset id", "asset"),
        "location": ("location",),
        "findings": ("findings", "finding", "evidence"),
        "recommendation": ("recommendation", "recommend"),
        "actions": ("actions", "action"),
        "restrictions": ("restrictions", "restriction"),
        "monitoring": ("monitoring", "monitor"),
    }
    out: dict[str, str] = {k: "" for k in keys}
    if not text or not text.strip():
        return out
    current: str | None = None
    buf: list[str] = []

    def _flush() -> None:
        nonlocal current, buf
        if current and current in out:
            out[current] = "\n".join(buf).strip()
        buf = []

    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.strip().lstrip("#").strip()
        lower = line.lower()
        matched_key = None
        rest = ""
        for key, names in aliases.items():
            for name in names:
                for prefix in (f"{name}:", f"**{name}**:"):
                    if lower.startswith(prefix):
                        matched_key = key
                        rest = line[len(prefix) :].strip()
                        break
                if matched_key:
                    break
            if matched_key:
                break
        if matched_key:
            _flush()
            current = matched_key
            buf = [rest] if rest else []
            continue
        if current:
            buf.append(raw.rstrip())
    _flush()
    return out


def polish_prompt(*, extract: str, cites: list[str], task_type: str) -> list[dict[str, str]]:
    """Messages for Gateway polish — structured fields only; no invented cites/numbers."""
    cite_block = "\n".join(f"- {c}" for c in cites) if cites else "- NOT FOUND"
    system = (
        "You polish industrial inspection notes for a DRAFT Word approval note.\n"
        "Rules: Use only facts in EXTRACT and CITES. Never invent measurements, codes, or citations.\n"
        "If a field is unknown write NOT STATED (never leave a field blank).\n"
        "Citations stay as provided — do not rewrite paths.\n"
        "decision must be one clear sentence (e.g. below-minimum thickness → recommend action).\n"
        "asset_id must capture any tag like V-101 if present in EXTRACT.\n"
        "Output ONLY these lines (plain text, no markdown fences):\n"
        "title: ...\n"
        "decision: ...\n"
        "asset_id: ...\n"
        "location: ...\n"
        "findings: ...\n"
        "recommendation: ...\n"
        "actions: ...\n"
        "restrictions: ...\n"
        "monitoring: ...\n"
    )
    user = (
        f"task_type={task_type}\n"
        f"--- EXTRACT ---\n{extract.strip()}\n"
        f"--- CITES ---\n{cite_block}\n"
        "Polish into the required fields."
    )
    return [
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ]


def extract_docx_text(path: Path) -> str:
    """Best-effort plain text from docx for secrets scan on export leave."""
    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)
    return "\n".join(parts)
