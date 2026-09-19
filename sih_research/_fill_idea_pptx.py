"""Fill SIH26117 working idea deck from official 2026 template. Do not overwrite the blank official file."""
from shutil import copyfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

SRC = r"C:\Users\THARUN PARSA\Documents\SIH26\sih_research\judging_sources\SIH2026-IDEA-Presentation-Format.pptx"
DST = r"C:\Users\THARUN PARSA\Documents\SIH26\sih_research\SIH26117_IDEA_WORKING.pptx"


def set_text(tf, text, size_pt, bold_first=False):
    tf.clear()
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size_pt)
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        if i == 0 and bold_first:
            run.font.bold = True


BODIES = {
    2: (
        "Proposed Solution (Describe your Idea/Solution/Prototype)\n"
        "• Sovereign Workbench — air-gapped agentic workbench on the org GPU (not a chatbot)\n"
        "• Scan → on-device OCR → task-routed open-weight models → HITL fact-sheet → Word + Excel\n"
        "• Coding task run and verified in Docker --network=none (GPU server stays outside the jail)\n"
        "\n"
        "How it addresses the problem\n"
        "• Confidential notes/P&IDs cannot go to Claude/Codex; PS requires own GPU + zero external calls\n"
        "• Auto-select ≥2 task types (extract vs code) via Model Cards + OpenAI-compatible /v1 gateway\n"
        "• Real deliverable (python-docx), citations or NOT FOUND, draft watermark, visible WAN-off log\n"
        "\n"
        "Innovation and uniqueness of the solution\n"
        "• vs Open WebUI/Ollama/PrivateGPT: task gateway, controlled templates, sandbox calc, HITL, egress proof\n"
        "• vs ChatGPT/Copilot: nothing leaves premises; 7B-class sequential on RTX 4060 8 GB (PS allows smaller models)\n"
        "• Not plant DCS control. Not a new LLM. Integration for industrial documents + proof of sovereignty",
        14,
    ),
    3: (
        "Technologies to be used (e.g. programming languages, frameworks, hardware)\n"
        "• Hardware: Laptop A RTX 4060 8 GB = runtime · Laptop B RTX 3050 6 GB = workbench UI (no LLM)\n"
        "• Software: Python workbench · OpenAI-compat /v1 (Ollama/vLLM/llama-server — pick at build) · CPU OCR\n"
        "• python-docx + openpyxl · Docker Desktop WSL2 --network=none for generated code + CR calc only\n"
        "• Local JSON KB (schema = future EAM) · on-disk Agent Skills (5) · audit log\n"
        "\n"
        "Methodology and process for implementation (Flow Charts/Images/ working prototype)\n"
        "  Engineer upload scan → CPU OCR → gateway vision_extract → JSON fact sheet → HITL approve\n"
        "       → sandbox CR=(t_prev-t_curr)/Δt → gateway draft_note (different model id) → .docx DRAFT\n"
        "  Parallel PS loop: prompt → code_sandbox → pytest in Docker jail → pass/fail shown\n"
        "  8 GB: unload/load one 7B-class at a time · WAN unplugged · never wrap GPU in the sandbox",
        13,
    ),
    4: (
        "Analysis of the feasibility of the idea\n"
        "• PS allows mid-range GPU / smaller models if 120B absent — 8 GB is that case\n"
        "• 36h REAL: OCR, two model ids, HITL JSON, Word, Excel CR, sandbox tests, WAN-off log\n"
        "• MOCKED (same schema): EAM/historian JSON · NOT claimed: SAP, 14B, PPT, two models in VRAM\n"
        "\n"
        "Potential challenges and risks\n"
        "• Small-model tool-calling / OCR on bad scans · Docker Desktop blocked at venue · fake-looking routing\n"
        "\n"
        "Strategies for overcoming these challenges\n"
        "• Sequential two GGUF ids (not two prompt labels) · OCR confidence + NOT FOUND · calc only in sandbox\n"
        "• Docker primary; WSL jail fallback (disclose if weaker) · WAN unplug + host/packet log · draft forever",
        14,
    ),
    5: (
        "Potential impact on the target audience\n"
        "• Primary: inspection / integrity engineer (MRPL beachhead)\n"
        "• Secondary: process, HSE, procurement, other PSU knowledge-work (same workbench, new templates later)\n"
        "\n"
        "Benefits of the solution (social, economic, environmental, etc.)\n"
        "• Time-to-DRAFT note: industry estimate scanned inspection 1–5 days → target hours (KB; not an MRPL SLA)\n"
        "• Cuts shadow ChatGPT on P&IDs / vendor / board data · audit trail · DPDP/CERT-In aligned logging (one line)\n"
        "• Economic: reuse existing GPU; no claimed crore savings · Environmental: no extra training run claimed\n"
        "• Scale: change base_url from laptop → plant GPU; HITL stays; never writes DCS/SIS/PLC",
        14,
    ),
    6: (
        "Details / Links of the reference and research work\n"
        "• SIH26117 Expected Solution (inspection→Word, sandbox coding, multimodal, zero-egress proof)\n"
        "• OISD-STD-128 (name/edition only) · API 510-class report fields (not pasted standard body)\n"
        "• OpenAI-compatible local serving docs (Ollama/vLLM/llama.cpp) · Agent Skills spec (on-disk)\n"
        "• Cursor sandboxing (code jail ≠ GPU) · SIH 2026 Guidelines: 6-slide PDF, nine criteria, no weights\n"
        "\n"
        "You vs existing (ticks only if true)     Air-gap  Citations  Sandbox  HITL  Word/Excel\n"
        "Sovereign Workbench (ours)                Y        Y          Y        Y     Y\n"
        "ChatGPT / Copilot / Claude                N        n/a        N        n/a   n/a\n"
        "Open WebUI / Ollama / PrivateGPT          local    N          N        weak  chat",
        13,
    ),
}


def main():
    copyfile(SRC, DST)
    prs = Presentation(DST)

    s1 = prs.slides[0]
    for sh in s1.shapes:
        if sh.name == "TextBox 9":
            set_text(
                sh.text_frame,
                "Problem Statement ID – SIH26117\n"
                "Problem Statement Title- Sovereign On-Premise Agentic AI Workbench\n"
                "using Open-Weight Multimodal LLMs for Confidential Industrial Work\n"
                "Theme- Smart Automation\n"
                "PS Category- Software\n"
                "Team ID- TBD (portal)\n"
                "Team Name (Registered on portal)- TBD",
                18,
            )

    s2 = prs.slides[1]
    for sh in s2.shapes:
        if sh.has_text_frame and (sh.text_frame.text or "").strip() == "IDEA TITLE":
            set_text(sh.text_frame, "Sovereign Workbench", 32, bold_first=True)

    for idx, (body, sz) in BODIES.items():
        slide = prs.slides[idx - 1]
        for sh in slide.shapes:
            if sh.name == "TextBox 8" and sh.has_text_frame:
                set_text(sh.text_frame, body, sz)
            if sh.has_text_frame and (sh.text_frame.text or "").strip() == "Your Team Name":
                set_text(sh.text_frame, "TBD", 12, bold_first=True)

    prs.save(DST)
    print("saved", DST, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
